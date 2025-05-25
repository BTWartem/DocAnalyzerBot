from pdfminer.high_level import extract_text
from docx import Document
from pptx import Presentation
import os

def parse_file(file_path: str) -> str:
    """Парсит PDF, DOCX, PPTX, TXT."""
    ext = os.path.splitext(file_path)[1].lower()
    text = ""

    try:
        if ext == ".pdf":
            text = extract_text(file_path)
        elif ext == ".docx":
            doc = Document(file_path)
            text = "\n".join([para.text for para in doc.paragraphs])
        elif ext == ".pptx":
            prs = Presentation(file_path)
            for slide in prs.slides:
                for shape in slide.shapes:
                    if hasattr(shape, "text"):
                        text += shape.text + "\n"
        elif ext == ".txt":
            with open(file_path, "r", encoding="utf-8") as f:
                text = f.read()
    except Exception as e:
        raise ValueError(f"Ошибка парсинга: {e}")

    return text.strip()